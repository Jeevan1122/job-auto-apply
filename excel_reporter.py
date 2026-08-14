from __future__ import annotations

"""
Generates a color-coded Excel report of all job applications.
Saved to the user's Desktop as: job_applications_YYYY-MM-DD.xlsx
"""
import json
import logging
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import EXCEL_OUTPUT_DIR
from database import get_all_jobs, get_all_daily_summaries

logger = logging.getLogger(__name__)

# ── Colors ────────────────────────────────────────────────────────────────────
COLOR_HEADER_BG   = "1F3864"   # dark navy
COLOR_HEADER_FONT = "FFFFFF"   # white
COLOR_APPLIED     = "C6EFCE"   # light green
COLOR_MANUAL      = "FFEB9C"   # light yellow
COLOR_FAILED      = "FFC7CE"   # light red/pink
COLOR_SKIPPED     = "EDEDED"   # light grey
COLOR_ROW_ALT     = "F2F7FF"   # very light blue (alternating rows)
COLOR_TITLE_BG    = "2E75B6"   # blue title banner


def _thin_border() -> Border:
    thin = Side(style="thin", color="D0D0D0")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_font()  -> Font:        return Font(bold=True, color=COLOR_HEADER_FONT, size=11)
def _title_font()   -> Font:        return Font(bold=True, color="FFFFFF", size=14)
def _normal_font()  -> Font:        return Font(size=10)
def _link_font()    -> Font:        return Font(size=10, color="0563C1", underline="single")
def _header_fill()  -> PatternFill: return PatternFill("solid", fgColor=COLOR_HEADER_BG)
def _applied_fill() -> PatternFill: return PatternFill("solid", fgColor=COLOR_APPLIED)
def _manual_fill()  -> PatternFill: return PatternFill("solid", fgColor=COLOR_MANUAL)
def _failed_fill()  -> PatternFill: return PatternFill("solid", fgColor=COLOR_FAILED)
def _skipped_fill() -> PatternFill: return PatternFill("solid", fgColor=COLOR_SKIPPED)
def _alt_fill()     -> PatternFill: return PatternFill("solid", fgColor=COLOR_ROW_ALT)


def _row_fill(status: str, row_idx: int) -> PatternFill:
    s = (status or "").lower()
    if s == "applied":          return _applied_fill()
    if s == "requires_manual":  return _manual_fill()
    if s == "failed":           return _failed_fill()
    if s in ("skipped", "scored", "found"): return _skipped_fill()
    return _alt_fill() if row_idx % 2 == 0 else PatternFill()


def _status_label(status: str) -> str:
    mapping = {
        "applied":          "✅ Applied",
        "requires_manual":  "⚠️ Manual Apply",
        "failed":           "❌ Failed",
        "skipped":          "⏭ Skipped",
        "scored":           "🔍 Found",
        "found":            "🔍 Found",
    }
    return mapping.get((status or "").lower(), status or "—")


# ── Applications sheet ─────────────────────────────────────────────────────────

APPS_COLUMNS = [
    ("Date",             15),
    ("Company",          22),
    ("Job Title",        35),
    ("Location",         20),
    ("Job Type",         13),
    ("Source",           15),
    ("Match %",          10),
    ("Applied?",         11),
    ("Resume Saved As",  30),
    ("Date Applied",     18),
    ("Job URL",          45),
    ("Notes",            30),
]


def _build_applications_sheet(ws, jobs: list[dict], today: str):
    # ── Title banner ───────────────────────────────────────────────────────────
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = f"Job Applications — {today}"
    title_cell.font  = _title_font()
    title_cell.fill  = PatternFill("solid", fgColor=COLOR_TITLE_BG)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── Legend ─────────────────────────────────────────────────────────────────
    ws.merge_cells("A2:L2")
    legend = ws["A2"]
    legend.value = (
        "  ✅ Applied (auto)      ⚠️ Manual Apply needed      "
        "❌ Failed      ⏭ Skipped (low match)"
    )
    legend.font = Font(size=9, italic=True, color="555555")
    legend.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── Header row ─────────────────────────────────────────────────────────────
    header_row = 3
    for col_idx, (col_name, col_width) in enumerate(APPS_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font      = _header_font()
        cell.fill      = _header_fill()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = _thin_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws.row_dimensions[header_row].height = 22

    # ── Data rows ──────────────────────────────────────────────────────────────
    today_jobs  = [j for j in jobs if j.get("date") == today]
    other_jobs  = [j for j in jobs if j.get("date") != today]
    sorted_jobs = today_jobs + other_jobs   # today first

    for row_idx, job in enumerate(sorted_jobs, start=header_row + 1):
        status  = job.get("status", "")
        score   = job.get("relevance_score") or 0.0
        fill    = _row_fill(status, row_idx)
        border  = _thin_border()
        applied_at = job.get("applied_at") or ""
        if applied_at and "T" in applied_at:
            try:
                applied_at = datetime.fromisoformat(applied_at).strftime("%Y-%m-%d %H:%M")
            except Exception:
                pass

        notes_raw = job.get("notes") or ""

        manually_applied = bool(job.get("manually_applied", 0))
        resume_filename  = job.get("resume_filename") or ""

        row_data = [
            job.get("date", ""),
            job.get("company", ""),
            job.get("title", ""),
            job.get("location", ""),
            (job.get("job_type") or "").title(),
            job.get("source", ""),
            f"{score:.0%}" if score else "—",
            "Yes" if manually_applied else "No",
            resume_filename,
            applied_at,
            job.get("url", ""),
            notes_raw if not notes_raw.startswith("[") else "",
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill      = fill
            cell.border    = border
            cell.alignment = Alignment(vertical="center", wrap_text=False)

            if col_idx == 7:   # Match % — center
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx == 8:   # Applied? — center
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx == 11 and value:  # URL — hyperlink style
                cell.font = _link_font()
                cell.hyperlink = value
                cell.value = "Open Job"

        ws.row_dimensions[row_idx].height = 18

    # ── Freeze panes & filters ─────────────────────────────────────────────────
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(APPS_COLUMNS))}{header_row}"
    )


# ── Summary sheet ──────────────────────────────────────────────────────────────

def _build_summary_sheet(ws, today: str):
    summaries = get_all_daily_summaries()

    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = "Daily Summary"
    t.font  = _title_font()
    t.fill  = PatternFill("solid", fgColor=COLOR_TITLE_BG)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["Date", "Jobs Found", "Applied", "Failed", "Manual Review", "Skipped", "Sheets Synced"]
    widths  = [14, 13, 11, 11, 15, 11, 14]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font      = _header_font()
        cell.fill      = _header_fill()
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _thin_border()
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[2].height = 20

    for ri, s in enumerate(summaries, start=3):
        is_today = s.get("date") == today
        fill = PatternFill("solid", fgColor="DDEEFF") if is_today else (
            _alt_fill() if ri % 2 == 0 else PatternFill()
        )
        row_vals = [
            s.get("date", ""),
            s.get("jobs_found", 0),
            s.get("jobs_applied", 0),
            s.get("jobs_failed", 0),
            s.get("jobs_manual", 0),
            s.get("jobs_skipped", 0),
            "✅" if s.get("sheets_updated") else "—",
        ]
        for ci, v in enumerate(row_vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill      = fill
            cell.border    = _thin_border()
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[ri].height = 18

    ws.freeze_panes = "A3"


# ── Main export ────────────────────────────────────────────────────────────────

def export_excel(today: str = None) -> str:
    if today is None:
        today = date.today().isoformat()

    jobs = get_all_jobs(limit=2000)

    wb = Workbook()

    # Sheet 1 — Applications
    ws_apps = wb.active
    ws_apps.title = "Applications"
    _build_applications_sheet(ws_apps, jobs, today)

    # Sheet 2 — Daily Summary
    ws_sum = wb.create_sheet("Daily Summary")
    _build_summary_sheet(ws_sum, today)

    # Save to Desktop
    out_dir  = Path(EXCEL_OUTPUT_DIR)
    out_path = out_dir / f"job_applications_{today}.xlsx"
    wb.save(str(out_path))

    logger.info("Excel report saved: %s", out_path)
    return str(out_path)
